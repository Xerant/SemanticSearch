from datetime import datetime
from pathlib import Path

import pytest
from _pdf import make_pdf

from semanticsearch.core import Element, ElementType, ParseError
from semanticsearch.ingest import SectionRules, clean_text, join_lines, normalize_elements
from semanticsearch.ingest.parsers import (
    CsvParser,
    DocxParser,
    EmlParser,
    PdfParser,
    TextParser,
    XlsxParser,
    default_parsers,
    find_parser,
)
from semanticsearch.packs.manifest import SectioningConfig

H, P, L, T, O = (
    ElementType.HEADING,
    ElementType.PARAGRAPH,
    ElementType.LIST_ITEM,
    ElementType.TABLE,
    ElementType.OTHER,
)
NUMBERED = SectionRules.from_config(SectioningConfig(heading_regexes=(r"^\d+(\.\d+)*\s+\S",)))
BASE = SectionRules()


def shape(elements):
    return [(e.type, e.text, e.level) for e in elements]


# ------------------------------------------------------------------ sectioning


@pytest.mark.parametrize(
    "text,level",
    [
        ("1 Introduction", 1),
        ("2.3 Risk Adjustments", 2),
        ("4.1.2 Settlement", 3),
        ("3.2 The participant must notify AEMO.", None),  # a clause, not a heading
        ("1 " + "x" * 200, None),  # too long
        ("2024 was a year", 1),  # regexes are the pack's responsibility; this one is loose
        ("Introduction", None),
    ],
)
def test_regex_level(text, level):
    assert NUMBERED.regex_level(text) == level


def test_no_regexes_means_no_regex_headings():
    assert BASE.regex_level("1 Introduction") is None


def test_style_level_default_and_custom():
    assert BASE.style_level("Heading 2") == 2
    assert BASE.style_level("Title") == 1
    assert BASE.style_level("Normal") is None
    custom = SectionRules(docx_styles=("Clause Title", "Clause Sub", "Heading 3"))
    assert custom.style_level("Clause Title") == 1
    assert custom.style_level("Clause Sub") == 2
    assert custom.style_level("Heading 3") == 3
    assert custom.style_level("Heading 1") is None  # explicit list replaces the defaults


# ------------------------------------------------------------------- normalise


def test_clean_text():
    assert clean_text("eﬃcient co­operation ​") == "efficient cooperation"
    assert clean_text("a\r\n\r\n  b  \t c") == "a\nb c"
    assert clean_text("5 m² and ½") == "5 m² and ½"  # NFKC would fold these


def test_join_lines_dehyphenates_only_lowercase_continuations():
    assert join_lines(["the contrac-", "tion of", "risk-", "Based pricing"]) == (
        "the contraction of risk- Based pricing"
    )


def test_normalize_elements():
    out = normalize_elements(
        [
            Element(P, "2.3 Risk\nAdjustments"),
            Element(P, "   "),
            Element(P, "Body text that is long enough\nto wrap."),
            Element(T, "a | b\nc | d"),
            Element(L, "1 item that looks numbered"),
        ],
        NUMBERED,
    )
    assert shape(out) == [
        (H, "2.3 Risk Adjustments", 2),
        (P, "Body text that is long enough to wrap.", 0),
        (T, "a | b\nc | d", 0),
        (L, "1 item that looks numbered", 0),  # only paragraphs are promoted
    ]


# ----------------------------------------------------------------------- text


def test_markdown(tmp_path):
    f = tmp_path / "a.md"
    f.write_text(
        "# Title\n\nIntro line one\nline two\n\n## Sub ##\n- first\n- second\n\n```\ncode\n```\n",
        encoding="utf-8",
    )
    assert shape(TextParser().parse(f)) == [
        (H, "Title", 1),
        (P, "Intro line one\nline two", 0),
        (H, "Sub", 2),
        (L, "first", 0),
        (L, "second", 0),
        (O, "code", 0),
    ]


def test_markdown_front_matter_breaks_and_tables(tmp_path):
    f = tmp_path / "a.md"
    f.write_text(
        "---\ntitle: Map\n---\n\n# Map\n\n---\n\n| Component | Evidence |\n|---|:---:|\n"
        "| Qdrant | [[Graph DB]] |\n| Ollama | x |\nAfter.\n\n***\n",
        encoding="utf-8",
    )
    assert shape(TextParser().parse(f)) == [
        (H, "Map", 1),
        (T, "Component | Evidence\nQdrant | [[Graph DB]]\nOllama | x", 0),
        (P, "After.", 0),
    ]


def test_plain_text_ignores_markdown_and_falls_back_on_encoding(tmp_path):
    f = tmp_path / "a.txt"
    f.write_bytes("# not a heading\n\ncaf\xe9 \x93quoted\x94".encode("latin-1"))
    assert shape(TextParser().parse(f)) == [(P, "# not a heading", 0), (P, "café “quoted”", 0)]


# ----------------------------------------------------------------------- docx


def test_docx(tmp_path):
    import docx

    d = docx.Document()
    d.add_paragraph("Hedge Book Review", style="Title")
    d.add_heading("Scope", level=1)
    d.add_paragraph("Covers cap contracts.")
    d.add_paragraph("   ")
    d.add_paragraph("first point", style="List Bullet")
    d.add_heading("Detail", level=2)
    t = d.add_table(rows=2, cols=3)
    t.cell(0, 0).merge(t.cell(0, 1)).text = "Merged"
    t.cell(0, 2).text = "C"
    for i, v in enumerate(("1", "2", "3")):
        t.cell(1, i).text = v
    d.add_paragraph("After table.")
    f = tmp_path / "a.docx"
    d.save(f)

    els = DocxParser(BASE).parse(f)
    assert shape(els) == [
        (H, "Hedge Book Review", 1),
        (H, "Scope", 1),
        (P, "Covers cap contracts.", 0),
        (L, "first point", 0),
        (H, "Detail", 2),
        (T, "Merged | C\n1 | 2 | 3", 0),
        (P, "After table.", 0),
    ]
    assert els[1].style == "Heading 1"
    assert all(e.page is None for e in els)


def test_docx_rejects_garbage(tmp_path):
    f = tmp_path / "bad.docx"
    f.write_bytes(b"not a zip")
    with pytest.raises(ParseError):
        DocxParser(BASE).parse(f)


# ------------------------------------------------------------------- tabular


def test_xlsx(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Trades"
    ws.append(["Trade", "Volume MW", "Date"])
    ws.append([None, None, None])
    ws.append(["Cap", 50.0, datetime(2025, 3, 1)])
    ws.append(["Swap", 12.5, None])
    wb.create_sheet("Empty")
    f = tmp_path / "a.xlsx"
    wb.save(f)

    assert shape(XlsxParser().parse(f)) == [
        (H, "Trades", 1),
        (T, "Trade: Cap | Volume MW: 50 | Date: 2025-03-01", 0),
        (T, "Trade: Swap | Volume MW: 12.5", 0),
    ]


def test_csv_sniffs_delimiter(tmp_path):
    f = tmp_path / "a.csv"
    f.write_text("region;price\nNSW1;101.5\nVIC1;;extra\n", encoding="utf-8")
    assert shape(CsvParser().parse(f)) == [
        (T, "region: NSW1 | price: 101.5", 0),
        (T, "region: VIC1 | column 3: extra", 0),
    ]


# ---------------------------------------------------------------------- email


def test_eml_plain(tmp_path):
    f = tmp_path / "a.eml"
    f.write_text(
        "From: a@example.com\nTo: b@example.com\nSubject: Cap settlement\n"
        "Content-Type: text/plain; charset=utf-8\n\nFirst para\nwraps.\n\n \nSecond.\n",
        encoding="utf-8",
    )
    assert shape(EmlParser().parse(f)) == [
        (H, "Cap settlement", 1),
        (O, "From: a@example.com\nTo: b@example.com", 0),
        (P, "First para\nwraps.", 0),
        (P, "Second.\n", 0),
    ]


def test_eml_html_only(tmp_path):
    f = tmp_path / "a.eml"
    f.write_text(
        "Subject: S\nContent-Type: text/html; charset=utf-8\n\n"
        "<html><head><style>p{}</style></head><body><p>One &amp; two</p><p>Three</p></body></html>\n",
        encoding="utf-8",
    )
    texts = [e.text.strip() for e in EmlParser().parse(f) if e.type is P]
    assert texts == ["One & two", "Three"]


# ------------------------------------------------------------------------ pdf


def report_pdf(tmp_path: Path) -> Path:
    def page(n, body):
        return [
            (72, 30, 9, "ACME Energy Quarterly Report"),  # running header
            *body,
            (280, 770, 9, f"Page {n} of 3"),  # running footer
        ]

    pages = [
        page(
            1,
            [
                (72, 100, 18, "Introduction"),  # font heading, largest -> level 1
                (72, 140, 10, "This report covers the hedge book and its contrac-"),
                (72, 152, 10, "tion over the quarter."),
                (72, 180, 10, "A second paragraph after a gap."),
                (72, 210, 14, "1.1 Scope"),  # regex heading; numbering sets the level
                (72, 240, 10, "- first bullet item"),
                (72, 252, 10, "(a) enumerated item"),
            ],
        ),
        page(2, [(72, 100, 10, "Continuation on page two."), (72, 112, 10, "Still the same paragraph.")]),
        page(
            3,
            [
                (72, 100, 18, "Appendix With A Very Long Title That"),
                (72, 120, 18, "Wraps Onto Two Lines"),
                (72, 160, 10, "Closing words."),
            ],
        ),
    ]
    f = tmp_path / "report.pdf"
    f.write_bytes(make_pdf(pages))
    return f


def test_pdf_structure_pages_and_running_lines(tmp_path):
    raw = PdfParser(NUMBERED).parse(report_pdf(tmp_path))
    els = normalize_elements(raw, NUMBERED)
    assert [(e.type, e.text, e.level, e.page) for e in els] == [
        (H, "Introduction", 1, 1),
        (P, "This report covers the hedge book and its contraction over the quarter.", 0, 1),
        (P, "A second paragraph after a gap.", 0, 1),
        (H, "1.1 Scope", 2, 1),
        (L, "first bullet item", 0, 1),
        (L, "(a) enumerated item", 0, 1),
        (P, "Continuation on page two. Still the same paragraph.", 0, 2),
        (H, "Appendix With A Very Long Title That Wraps Onto Two Lines", 1, 3),
        (P, "Closing words.", 0, 3),
    ]


def test_pdf_heuristic_none_uses_regexes_only(tmp_path):
    rules = SectionRules.from_config(
        SectioningConfig(heading_regexes=(r"^\d+(\.\d+)*\s+\S",), pdf_heuristic="none")
    )
    headings = [e.text for e in PdfParser(rules).parse(report_pdf(tmp_path)) if e.type is H]
    assert headings == ["1.1 Scope"]


def test_pdf_without_text_is_empty(tmp_path):
    f = tmp_path / "scan.pdf"
    f.write_bytes(make_pdf([[]]))
    assert PdfParser(BASE).parse(f) == []


def test_pdf_rejects_garbage(tmp_path):
    f = tmp_path / "bad.pdf"
    f.write_bytes(b"%PDF-1.4 truncated")
    with pytest.raises(ParseError):
        PdfParser(BASE).parse(f)


# -------------------------------------------------------------------- registry


@pytest.mark.parametrize(
    "name,cls",
    [
        ("a.PDF", PdfParser),
        ("a.docx", DocxParser),
        ("a.xlsx", XlsxParser),
        ("a.csv", CsvParser),
        ("a.eml", EmlParser),
        ("a.md", TextParser),
        ("a.doc", type(None)),
    ],
)
def test_find_parser(name, cls):
    assert type(find_parser(default_parsers(BASE), Path(name))) is cls
